"""Черновик отбора РОССКО → Авито, Иркутск. Не публикует объявления."""
from __future__ import annotations

import argparse
import csv
import io
import re
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook

from update_prices import (get_cell, load_config, normalize_key, parse_decimal,
                           parse_stock_count, read_header_map, read_photo_map, sale_price, text)

REQUIRED = ("Бренд", "Артикул", "Описание", "Цена, руб.", "Наличие", "Срок поставки, дн.")
CITY_FIELDS = ("Город склада", "Склад отгрузки", "Филиал", "Город")
EXCLUDE = re.compile(r"ремкомплект|ремкомпл|тяга|наконечник|пыльник|колодк|суппорт|диск сцепления|втулк|болт|подшипник|щиток", re.I)
RACK = re.compile(r"(?:рулев\w*\s+рейк\w*|рейк\w*\s+рулев\w*)", re.I)
BRAKE = re.compile(r"(?:тормозн\w*\s+диск\w*|диск\w*\s+тормозн\w*)", re.I)
FIELDS = ("Группа", "Бренд", "Артикул", "Описание РОССКО", "Цена Авито", "Остаток РОССКО",
          "Срок поставки", "Город", "Адрес выдачи", "Применимость", "Каталожный номер", "Фото для проверки", "Статус проверки")
REVIEW_FIELDS = ("brand", "article", "approved", "pickup_confirmed",
                 "condition_confirmed", "photo_urls", "fitment", "oem_number")


def category(name: str) -> str | None:
    s = re.sub(r"[-_/(),;]+", " ", name.casefold())
    if EXCLUDE.search(s):
        return None
    if RACK.search(s):
        return "Рулевые рейки"
    if BRAKE.search(s):
        return "Тормозные диски"
    return None


def zero_days(value: object) -> bool:
    return bool(re.fullmatch(r"0+(?:[.,]0+)?", text(value)))


def select_rows(ws, config: dict, photos: dict[str, str] | None = None) -> list[dict]:
    headers = read_header_map(ws)
    missing = set(REQUIRED) - set(headers)
    if missing:
        raise ValueError("В прайсе отсутствуют поля: " + ", ".join(sorted(missing)))
    cities = [x for x in CITY_FIELDS if x in headers]
    pickup_address = text(config.get("avito", {}).get("pickup_address"))
    if not pickup_address:
        raise ValueError("Не указан адрес выдачи Авито в настройках avito.pickup_address")
    seen: set[str] = set()
    result = []
    for row in ws.iter_rows(min_row=2):
        desc = text(get_cell(row, headers, "Описание"))
        group = category(desc)
        if not group:
            continue
        if cities and not any("иркутск" in text(get_cell(row, headers, h)).casefold() for h in cities):
            continue
        if not zero_days(get_cell(row, headers, "Срок поставки, дн.")):
            continue
        raw_stock = text(get_cell(row, headers, "Наличие"))
        # Не трактовать отрицательный остаток как положительный: общий парсер извлекает цифры.
        if raw_stock.lstrip().startswith("-"):
            continue
        quantity = parse_stock_count(raw_stock)
        price = parse_decimal(get_cell(row, headers, "Цена, руб."))
        brand = text(get_cell(row, headers, "Бренд"))
        article = text(get_cell(row, headers, "Артикул"))
        if not brand or not article or not price or price <= 0 or quantity <= 0:
            continue
        key = normalize_key(brand) + "|" + normalize_key(article)
        if key in seen:
            continue
        seen.add(key)
        result.append({
            "Группа": group, "Бренд": brand, "Артикул": article,
            "Описание РОССКО": desc, "Цена Авито": sale_price(price, config.get("pricing", {})),
            "Остаток РОССКО": quantity, "Срок поставки": 0, "Город": "Иркутск", "Адрес выдачи": pickup_address,
            "Применимость": text(get_cell(row, headers, "Применимость")),
            "Каталожный номер": text(get_cell(row, headers, "Каталожный номер")),
            "Фото для проверки": (photos or {}).get(key, ""),
            "Статус проверки": "Черновик — проверить фото, применимость и наличие",
        })
    return result


def load_fresh_price(config: dict, base: Path, max_age_hours: float):
    source = config.get("rossko", {})
    mode = text(source.get("price_source") or "local").casefold()
    now = datetime.now(timezone.utc)
    if mode == "url":
        url = text(source.get("price_url"))
        if not url.startswith("https://"):
            raise ValueError("Для РОССКО нужна HTTPS-ссылка")
        response = requests.get(url, timeout=int(source.get("download_timeout_seconds") or 120))
        response.raise_for_status()
        data = response.content
        if len(data) < 10000 or not data.startswith(b"PK"):
            raise ValueError("РОССКО вернул некорректный файл")
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if mode != "local":
        raise ValueError("Неизвестный источник прайса")
    path = base / text(source.get("local_price_path") or "input/rossko_price.xlsx")
    if not path.exists():
        raise FileNotFoundError(path)
    age = (now.timestamp() - path.stat().st_mtime) / 3600
    if age < 0 or age > max_age_hours:
        raise ValueError("Локальный прайс старше допустимого срока; выгрузка остановлена")
    return load_workbook(path, read_only=True, data_only=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default="config.json")
    parser.add_argument("--confirm-irkutsk-source", action="store_true")
    parser.add_argument("--max-local-age-hours", type=float, default=2)
    parser.add_argument("--out", default="output/avito_irkutsk_candidates.csv")
    parser.add_argument("--pilot-per-category", type=int, default=25, help="Максимум кандидатов каждой группы (0 = все)")
    args = parser.parse_args()
    if not args.confirm_irkutsk_source:
        parser.error("Подтвердите, что исходный прайс относится к складу РОССКО в Иркутске")
    if args.max_local_age_hours <= 0:
        parser.error("Срок свежести должен быть положительным")
    config_path = Path(args.config).resolve()
    cfg = load_config(config_path)
    if args.pilot_per_category < 0:
        parser.error("--pilot-per-category не может быть отрицательным")
    photos_path = config_path.parent / text(cfg.get("drom", {}).get("photo_map_path") or "input/photos.csv")
    photos = read_photo_map(photos_path)
    wb = load_fresh_price(cfg, config_path.parent, args.max_local_age_hours)
    try:
        rows = select_rows(wb.active, cfg, photos)
    finally:
        wb.close()
    if not rows:
        print("Нет товаров, удовлетворяющих фильтру. Публикация запрещена.")
        return 2
    total_counts = {group: sum(p["Группа"] == group for p in rows) for group in ("Рулевые рейки", "Тормозные диски")}
    if args.pilot_per_category:
        pilot = []
        for group in ("Рулевые рейки", "Тормозные диски"):
            sorted_group = sorted((p for p in rows if p["Группа"] == group),
                                  key=lambda p: (-p["Остаток РОССКО"], p["Бренд"], p["Артикул"]))
            pilot.extend(sorted_group[:args.pilot_per_category])
        rows = pilot
    target = config_path.parent / args.out
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS + REVIEW_FIELDS, delimiter=";")
        writer.writeheader()
        for row in rows:
            review = dict(row)
            review.update({"brand": row["Бренд"], "article": row["Артикул"],
                           "approved": "", "pickup_confirmed": "",
                           "condition_confirmed": "", "photo_urls": "",
                           "fitment": "", "oem_number": ""})
            writer.writerow(review)
    print(f"Свежий прайс: найдено всего по категориям {total_counts}; отобрано для проверки: {len(rows)}; с предлагаемым фото: {sum(bool(p['Фото для проверки']) for p in rows)}")
    print(f"Черновик: {target}")
    print("Это НЕ файл автозагрузки. Объявления не созданы.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
