#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Rossko -> Drom auto-link exporter.

Generates stable XLSX price files for Drom from a Rossko daily XLSX price list.
The output filename stays the same, so Drom can pull it by a permanent URL.

No credentials are embedded here. If Rossko provides a direct price URL, put it
into config.json. Otherwise place the latest rossko_price.xlsx into input/.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import re
import shutil
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, ROUND_CEILING, InvalidOperation
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import requests
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension

DEFAULT_DROM_HEADERS = [
    "Артикул", "Наименование товара", "Новый/б.у.", "Марка", "Модель", "Кузов", "Номер",
    "Производитель", "Двигатель", "Год", "LR", "FR", "UD", "Цвет", "Номера замен", "Примеч.",
    "Кол-во", "Цена", "Наличие", "Сроки доставки", "Фотография", "Поставщик", "ИНН поставщика", "Адрес склада",
]

REQUIRED_ROSSKO_HEADERS = ["Бренд", "Артикул", "Описание", "Цена, руб.", "Наличие", "Срок поставки, дн."]


@dataclass(frozen=True)
class Product:
    brand: str
    article: str
    description: str
    price: Decimal
    stock_count: int
    delivery_days: int
    catalog_number: str = ""
    oem_number: str = ""
    application: str = ""
    vendor_code: str = ""


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_key(value: Any) -> str:
    return re.sub(r"[^0-9a-zа-я]+", "", text(value).casefold())


def parse_decimal(value: Any) -> Optional[Decimal]:
    raw = text(value).replace(" ", "").replace(",", ".")
    if not raw:
        return None
    try:
        return Decimal(raw)
    except (InvalidOperation, ValueError):
        return None


def parse_stock_count(value: Any) -> int:
    """Drom should not overstate stock. For ranges like 10-100, take 10."""
    raw = text(value).replace(" ", "")
    if not raw:
        return 0
    raw = raw.replace(",", ".")
    nums = re.findall(r"\d+(?:\.\d+)?", raw)
    if not nums:
        return 0
    return int(float(nums[0]))


def parse_int(value: Any, default: int = 0) -> int:
    raw = text(value).replace(" ", "").replace(",", ".")
    if not raw:
        return default
    try:
        return int(float(raw))
    except ValueError:
        nums = re.findall(r"\d+", raw)
        return int(nums[0]) if nums else default


def load_config(path: Path) -> Dict[str, Any]:
    if not path.exists():
        example = path.with_name("config.example.json")
        if example.exists():
            shutil.copyfile(example, path)
            print(f"Создан {path.name} из config.example.json. Проверь настройки и запусти повторно.")
            sys.exit(2)
        raise FileNotFoundError(f"Не найден config: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def download_price_if_needed(config: Dict[str, Any], base_dir: Path) -> Path:
    rossko = config.get("rossko", {})
    source_type = text(rossko.get("price_source") or "local").lower()
    local_path = base_dir / text(rossko.get("local_price_path") or "input/rossko_price.xlsx")

    if source_type == "url":
        url = text(rossko.get("price_url"))
        if not url:
            raise ValueError("В config.json выбран price_source=url, но price_url пустой.")
        timeout = int(rossko.get("download_timeout_seconds") or 120)
        local_path.parent.mkdir(parents=True, exist_ok=True)
        tmp_path = local_path.with_suffix(".download")
        print("Скачиваю прайс Росско по ссылке...")
        response = requests.get(url, timeout=timeout)
        response.raise_for_status()

        # Простая проверка: не перезаписывать рабочий файл HTML-страницей ошибки
        # или пустым ответом. Хэши/сравнение версий специально не используем.
        content = response.content
        if len(content) < 10_000 or not content.startswith(b"PK"):
            raise ValueError("Росско отдал не XLSX-файл или файл слишком маленький. Старый прайс не перезаписан.")

        tmp_path.write_bytes(content)
        tmp_path.replace(local_path)
        print(f"Прайс обновлен: {local_path}")

    if not local_path.exists():
        raise FileNotFoundError(f"Не найден прайс Росско: {local_path}")
    return local_path


def read_header_map(ws) -> Dict[str, int]:
    header_map: Dict[str, int] = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        name = text(cell.value)
        if name:
            header_map[name] = col_idx
    return header_map


def get_cell(row, header_map: Dict[str, int], header: str) -> Any:
    idx = header_map.get(header)
    if not idx:
        return None
    return row[idx - 1].value


def read_rossko_products(path: Path, config: Dict[str, Any]) -> List[Product]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    header_map = read_header_map(ws)
    missing = [h for h in REQUIRED_ROSSKO_HEADERS if h not in header_map]
    if missing:
        raise ValueError("В прайсе Росско не найдены колонки: " + ", ".join(missing))

    filters = config.get("filters", {})
    skip_without_stock = bool(filters.get("skip_without_stock", True))
    skip_without_price = bool(filters.get("skip_without_price", True))
    max_delivery_days = filters.get("max_delivery_days")
    allowed_brands = {normalize_key(x) for x in filters.get("allowed_brands", []) if text(x)}
    blocked_brands = {normalize_key(x) for x in filters.get("blocked_brands", []) if text(x)}
    blocked_articles = {normalize_key(x) for x in filters.get("blocked_articles", []) if text(x)}

    products: List[Product] = []
    seen: set[str] = set()
    skipped = {"duplicates": 0, "no_stock": 0, "no_price": 0, "filtered": 0}

    for row in ws.iter_rows(min_row=2):
        brand = text(get_cell(row, header_map, "Бренд"))
        article = text(get_cell(row, header_map, "Артикул"))
        description = text(get_cell(row, header_map, "Описание"))
        if not brand or not article or not description:
            skipped["filtered"] += 1
            continue

        brand_key = normalize_key(brand)
        article_key = normalize_key(article)
        identity = f"{brand_key}|{article_key}"
        if identity in seen:
            skipped["duplicates"] += 1
            continue
        if allowed_brands and brand_key not in allowed_brands:
            skipped["filtered"] += 1
            continue
        if brand_key in blocked_brands or article_key in blocked_articles:
            skipped["filtered"] += 1
            continue

        price = parse_decimal(get_cell(row, header_map, "Цена, руб."))
        if price is None or price <= 0:
            if skip_without_price:
                skipped["no_price"] += 1
                continue
            price = Decimal("0")

        stock_count = parse_stock_count(get_cell(row, header_map, "Наличие"))
        if stock_count <= 0 and skip_without_stock:
            skipped["no_stock"] += 1
            continue

        delivery_days = max(parse_int(get_cell(row, header_map, "Срок поставки, дн."), 0), 0)
        if max_delivery_days is not None and delivery_days > int(max_delivery_days):
            skipped["filtered"] += 1
            continue

        seen.add(identity)
        products.append(Product(
            brand=brand,
            article=article,
            description=description,
            price=price,
            stock_count=max(stock_count, 0),
            delivery_days=delivery_days,
            catalog_number=text(get_cell(row, header_map, "Каталожный номер")),
            oem_number=text(get_cell(row, header_map, "OEМ Номер")),
            application=text(get_cell(row, header_map, "Применимость")),
            vendor_code=text(get_cell(row, header_map, "Вендор-код")),
        ))

    print(f"Товаров к выгрузке: {len(products)}")
    print("Пропущено: " + ", ".join(f"{k}={v}" for k, v in skipped.items() if v))
    return products


def read_drom_headers(template_path: Path) -> List[str]:
    if not template_path.exists():
        print("Шаблон Дрома не найден, использую стандартные заголовки.")
        return DEFAULT_DROM_HEADERS
    wb = load_workbook(template_path, read_only=True, data_only=True)
    ws = wb.active
    headers = [text(cell.value) for cell in ws[1] if text(cell.value)]
    return headers or DEFAULT_DROM_HEADERS


def read_photo_map(path: Path) -> Dict[str, str]:
    photos: Dict[str, str] = {}
    if not path.exists():
        return photos
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.DictReader(f, delimiter=delimiter)
        for row in reader:
            brand = row.get("brand") or row.get("Бренд") or row.get("производитель") or ""
            article = row.get("article") or row.get("Артикул") or row.get("номер") or ""
            photo = row.get("photo_url") or row.get("Фотография") or row.get("url") or ""
            if text(brand) and text(article) and text(photo):
                photos[f"{normalize_key(brand)}|{normalize_key(article)}"] = text(photo)
    return photos


def sale_price(price: Decimal, pricing: Dict[str, Any]) -> int:
    markup_percent = Decimal(str(pricing.get("markup_percent", 37)))
    round_up_to = Decimal(str(pricing.get("round_up_to", 50)))
    min_price = Decimal(str(pricing.get("min_price", 0)))
    result = price * (Decimal("1") + markup_percent / Decimal("100"))
    if min_price > 0 and result < min_price:
        result = min_price
    if round_up_to > 0:
        result = (result / round_up_to).to_integral_value(rounding=ROUND_CEILING) * round_up_to
    return int(result)


def delivery_text(product: Product, profile: Dict[str, Any]) -> str:
    availability_text = text(profile.get("availability_text") or "Под заказ")
    local_delivery_text = text(profile.get("local_delivery_text") or "2-5 дней")
    extra_min = int(profile.get("extra_days_min") or 0)
    extra_max = int(profile.get("extra_days_max") or 0)
    use_source = bool(profile.get("use_source_delivery_days", True))

    # In a Rossko price formed for Irkutsk, 0 days means immediate local availability for this direction.
    # For Drom the status remains "Под заказ"; only the delivery text changes.
    if product.delivery_days <= 0:
        return local_delivery_text
    if not use_source:
        return local_delivery_text
    if extra_min == 0 and extra_max == 0:
        return f"{product.delivery_days} дней"
    return f"{product.delivery_days + extra_min}-{product.delivery_days + extra_max} дней"


def drom_row(product: Product, headers: Sequence[str], config: Dict[str, Any], profile: Dict[str, Any], photos: Dict[str, str]) -> List[Any]:
    drom = config.get("drom", {})
    pricing = config.get("pricing", {})
    key = f"{normalize_key(product.brand)}|{normalize_key(product.article)}"
    fields: Dict[str, Any] = {
        "Артикул": product.article,
        "Наименование товара": product.description,
        "Новый/б.у.": "Новый",
        "Марка": "",
        "Модель": "",
        "Кузов": "",
        "Номер": product.catalog_number or product.article,
        "Производитель": product.brand,
        "Двигатель": "",
        "Год": "",
        "LR": "",
        "FR": "",
        "UD": "",
        "Цвет": "",
        "Номера замен": product.oem_number,
        "Примеч.": product.application,
        "Кол-во": product.stock_count,
        "Цена": sale_price(product.price, pricing),
        "Наличие": text(profile.get("availability_text") or "Под заказ"),
        "Сроки доставки": delivery_text(product, profile),
        "Фотография": photos.get(key, ""),
        "Поставщик": text(drom.get("supplier_name")),
        "ИНН поставщика": text(drom.get("supplier_inn")),
        "Адрес склада": text(drom.get("stock_address")),
    }
    return [fields.get(h, "") for h in headers]


def write_drom_xlsx(path: Path, headers: Sequence[str], rows: Iterable[Sequence[Any]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "Drom"
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    count = 0
    for row in rows:
        ws.append(list(row))
        count += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 18, "B": 48, "C": 14, "D": 16, "E": 18, "F": 16, "G": 18, "H": 18,
        "I": 14, "J": 12, "K": 8, "L": 8, "M": 8, "N": 12, "O": 30, "P": 55,
        "Q": 10, "R": 12, "S": 16, "T": 18, "U": 45, "V": 22, "W": 18, "X": 40,
    }
    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = widths.get(letter, 18)
    wb.save(path)
    return count


def write_report(path: Path, profile_name: str, output_file: str, exported: int, products_total: int) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    line = f"{now};{profile_name};{output_file};exported={exported};source_products={products_total}\n"
    with path.open("a", encoding="utf-8") as f:
        f.write(line)


def run_once(config_path: Path) -> int:
    base_dir = config_path.resolve().parent
    config = load_config(config_path)
    public_dir = base_dir / text(config.get("drom", {}).get("public_dir") or "public")
    template_path = base_dir / text(config.get("drom", {}).get("template_path") or "input/auto-parts-GT.xlsx")
    photo_map_path = base_dir / text(config.get("drom", {}).get("photo_map_path") or "input/photos.csv")

    price_path = download_price_if_needed(config, base_dir)
    headers = read_drom_headers(template_path)
    photos = read_photo_map(photo_map_path)
    products = read_rossko_products(price_path, config)

    enabled_profiles = [p for p in config.get("profiles", []) if p.get("enabled", True)]
    if not enabled_profiles:
        raise ValueError("В config.json нет включенных профилей города.")

    stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for profile in enabled_profiles:
        filename = text(profile.get("output_filename")) or f"drom_{normalize_key(profile.get('name'))}.xlsx"
        output_path = public_dir / filename
        rows = (drom_row(product, headers, config, profile, photos) for product in products)
        exported = write_drom_xlsx(output_path, headers, rows)
        write_report(base_dir / "logs" / "updates.log", text(profile.get("name")), filename, exported, len(products))
        print(f"[{stamp}] {profile.get('name')}: создан {output_path} строк={exported}")

    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Rossko -> Drom stable XLSX exporter")
    parser.add_argument("--config", default="config.json", help="Путь к config.json")
    parser.add_argument("--loop", type=int, default=0, help="Повторять обновление каждые N секунд")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    config_path = Path(args.config)
    if args.loop and args.loop > 0:
        while True:
            try:
                run_once(config_path)
            except Exception as exc:
                print(f"Ошибка обновления: {exc}", file=sys.stderr)
            time.sleep(args.loop)
    return run_once(config_path)


if __name__ == "__main__":
    raise SystemExit(main())
