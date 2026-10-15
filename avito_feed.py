"""Формирование двух Excel-файлов Авито из свежего прайса РОССКО.

По умолчанию безопасно: без подтверждений из CSV карточки не выгружаются.
Выпуск XLSX не публикует объявления автоматически.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import re
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook

from avito_irkutsk import load_fresh_price, select_rows
from update_prices import load_config, normalize_key, text

# Названия и порядок строго соответствуют двум шаблонам Авито от 13.10.2026.
HEADERS = (
    "Адрес", "Широта", "Долгота", "Уникальный идентификатор объявления",
    "Начало размещения", "Окончание размещения", "Способ размещения",
    "Услуга продвижения", "Номер объявления на Авито", "Контактное лицо",
    "Номер телефона", "Идентификатор адреса", "Способ связи", "Addresses",
    "Адреса отгрузки", "Категория", "Описание объявления", "Названия фото",
    "Ссылки на фото", "Ссылка на видео",
    "Настройка цены целевого действия",
    "Настройка цены целевого действия: автоматическая",
    "Настройка цены целевого действия: ручная",
    "Название объявления", "Интернет звонки", "Устройства для приёма звонков",
    "Способ доставки", "Вес (Для Доставки)", "Длина (Для Доставки)",
    "Высота (Для Доставки)", "Ширина (Для Доставки)", "Возвраты",
    "Субсидирование доставки", "ТН ВЭД", "ОКПД 2", "Документы не нужны",
    "Честный знак", "Коды GTIN", "Сертификат", "Декларация", "СГР",
    "СГР Дата", "СГР Ссылка", "Минсельхоз", "Минсельхоз Дата", "РУ",
    "РУ Дата", "ГИИС ДМДК", "Цена", "URL видеофайла",
    "Вид товара", "Вид объявления", "Тип товара", "Вид запчасти",
    "Состояние", "Происхождение", "Доступность", "Производитель",
    "Номер детали OEM", "Производитель оригинала", "Номер оригинала",
    "Авто для которых подходит запчасть", "Включая НДС",
    "Оптовые продажи", "Тип минимального заказа",
    "Количество в минимальном заказе", "Фасовка", "Количество в фасовке",
    "В чём измеряются товары", "Скидка за опт", "Размер скидки за опт",
)

GROUPS = {
    "Рулевые рейки": {
        "category": "Рулевое управление",
        "filename": "avito_irkutsk_steering.xlsx",
        "template_key": "template_steering",
        "name": "Рулевая рейка",
    },
    "Тормозные диски": {
        "category": "Тормозная система",
        "filename": "avito_irkutsk_brakes.xlsx",
        "template_key": "template_brakes",
        "name": "Тормозной диск",
    },
}
CATEGORY_PREFIX = "Транспорт - Запчасти и аксессуары - Запчасти - Для автомобилей - "
REQUIRED = {
    "Адрес", "Уникальный идентификатор объявления", "Категория",
    "Описание объявления", "Название объявления", "Цена", "Вид товара",
    "Вид объявления", "Тип товара", "Вид запчасти", "Состояние",
}
APPROVAL_FIELDS = [
    "brand", "article", "approved", "pickup_confirmed",
    "condition_confirmed", "photo_urls", "fitment", "oem_number",
]
TRUE_VALUES = {"да", "yes", "true", "1", "y"}


def is_yes(value: object) -> bool:
    return text(value).casefold() in TRUE_VALUES


def product_key(brand: str, article: str) -> str:
    return normalize_key(brand) + "|" + normalize_key(article)


def stable_id(group: str, brand: str, article: str) -> str:
    source = "|".join([group, normalize_key(brand), normalize_key(article)])
    return "irk-" + hashlib.sha256(source.encode("utf-8")).hexdigest()[:22]


def safe_image_urls(value: str) -> str:
    urls = [v.strip() for v in re.split(r"[|\n]+", value) if v.strip()]
    if not urls or len(urls) > 10:
        return ""
    for url in urls:
        parsed = urlparse(url)
        if parsed.scheme != "https" or not parsed.netloc or " " in url or len(url) > 2048:
            return ""
    return "|".join(urls)


def load_approvals(path: Path) -> dict[str, dict]:
    if not path.is_file():
        raise FileNotFoundError(
            f"Нет файла подтверждений: {path}. Создайте CSV по образцу input/avito_approved.example.csv."
        )
    approvals = {}
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, delimiter=";")
        missing = set(APPROVAL_FIELDS) - set(reader.fieldnames or [])
        if missing:
            raise ValueError("В CSV подтверждений отсутствуют: " + ", ".join(sorted(missing)))
        for row in reader:
            key = product_key(row["brand"], row["article"])
            if not key or key == "|":
                continue
            if key in approvals:
                raise ValueError("Дублирование товара в CSV подтверждений: " + key)
            approvals[key] = row
    return approvals


def approved_payload(product: dict, approval: dict | None) -> tuple[dict | None, str]:
    if approval is None or not is_yes(approval.get("approved")):
        return None, "нет подтверждения"
    if not is_yes(approval.get("pickup_confirmed")):
        return None, "не подтверждена выдача по адресу"
    if not is_yes(approval.get("condition_confirmed")):
        return None, "не подтверждено состояние «Новое»"
    fitment = text(approval.get("fitment"))
    if product["Группа"] == "Тормозные диски" and not fitment:
        return None, "не указана проверенная применимость к авто"
    urls = safe_image_urls(text(approval.get("photo_urls")))
    if not urls:
        return None, "нет проверенных HTTPS-фотографий"
    return {"fitment": fitment, "photo_urls": urls,
            "oem_number": text(approval.get("oem_number"))}, ""


def make_fields(product: dict, verified: dict, address: str) -> dict:
    group = GROUPS[product["Группа"]]
    brand = product["Бренд"]
    article = product["Артикул"]
    title = f"{group['name']} {brand} {article}".strip()
    if len(title) > 100:
        title = title[:100].rstrip()
    desc = (
        f"{group['name']}. Производитель: {brand}.\n"
        f"Артикул: {article}.\n"
        f"Состояние: новое.\n"
        f"Выдача: {address}.\n"
        "Наличие актуально на момент проверки прайса поставщика. "
        "Перед приездом уточните доступность и совместимость по VIN."
    )
    if verified["fitment"]:
        desc += "\nПодтверждённая применимость: " + verified["fitment"]
    return {
        "Адрес": address,
        "Уникальный идентификатор объявления": stable_id(product["Группа"], brand, article),
        "Категория": "Запчасти и аксессуары",
        "Описание объявления": desc,
        "Ссылки на фото": verified["photo_urls"],
        "Название объявления": title,
        "Цена": product["Цена Авито"],
        "Вид товара": "Запчасти",
        "Вид объявления": "Товар приобретен на продажу",
        "Тип товара": "Для автомобилей",
        "Вид запчасти": group["category"],
        "Состояние": "Новое",
        "Доступность": "В наличии",
        "Производитель": brand,
        "Номер детали OEM": verified["oem_number"],
        "Авто для которых подходит запчасть": verified["fitment"],
    }


def validate_row(fields: dict, group_name: str) -> None:
    required = REQUIRED | ({"Авто для которых подходит запчасть"} if group_name == "Тормозные диски" else set())
    empty = [f for f in sorted(required) if fields.get(f) in (None, "")]
    if empty:
        raise ValueError("Незаполненные обязательные поля: " + ", ".join(empty))


def make_book(fields_list: list[dict], group_name: str, template: Path | None = None):
    if template:
        if not template.is_file():
            raise FileNotFoundError(f"Не найден шаблон Авито: {template}")
        wb = load_workbook(template)
        # Avito may parse visible reference tabs as listings without an ID.
        # Hide the lookups, preserving their Excel dropdown references.
        for auxiliary in wb.worksheets:
            if auxiliary.title not in {"Инструкция", "Объявления"}:
                auxiliary.sheet_state = "hidden"
        if "Объявления" not in wb.sheetnames:
            raise ValueError("Шаблон Авито не содержит лист «Объявления»")
        ws = wb["Объявления"]
        actual = [text(ws.cell(2, col).value) for col in range(1, len(HEADERS)+1)]
        if tuple(actual) != HEADERS:
            raise ValueError("Шаблон Авито не совпадает с ожидаемыми 71 полями")
        if not text(ws.cell(1, 1).value).endswith(GROUPS[group_name]["category"]):
            raise ValueError("Выбран шаблон другой категории")
        if any(ws.cell(r, 1).value for r in range(5, ws.max_row + 1)):
            raise ValueError("Шаблон уже содержит объявления: необходим пустой оригинал")
    else:
        # Точная последовательность полей из присланных шаблонов; без справочников
        # это техническая копия и требует проверки в Авито перед публикацией.
        wb = Workbook()
        ws = wb.active
        ws.title = "Объявления"
        ws.cell(1, 1, CATEGORY_PREFIX + GROUPS[group_name]["category"])
        for col, header in enumerate(HEADERS, 1):
            ws.cell(2, col, header)
            ws.cell(3, col, "Обязательный" if header in REQUIRED else "Необязательный")
            ws.cell(4, col, "Подробнее о параметре")
        if group_name == "Тормозные диски":
            ws.cell(3, HEADERS.index("Авто для которых подходит запчасть") + 1, "Обязательный")
        ws.freeze_panes = "A5"
    for index, fields in enumerate(fields_list, 5):
        validate_row(fields, group_name)
        for col, header in enumerate(HEADERS, 1):
            value = fields.get(header)
            if value not in (None, ""):
                ws.cell(index, col, value)
    return wb


def build(config_path: Path, approvals_path: Path, output_dir: Path, *, allow_generated_template: bool = False):
    cfg = load_config(config_path)
    avito = cfg.get("avito", {})
    address = text(avito.get("pickup_address"))
    if not address or "иркутск" not in address.casefold():
        raise ValueError("Нужен подтверждённый адрес выдачи в Иркутске")
    wb = load_fresh_price(cfg, config_path.parent, 2)
    try:
        candidates = select_rows(wb.active, cfg)
    finally:
        wb.close()
    approvals = load_approvals(approvals_path)
    rows = {group: [] for group in GROUPS}
    skipped: dict[str, int] = {}
    for product in candidates:
        verified, reason = approved_payload(product, approvals.get(product_key(product["Бренд"], product["Артикул"])))
        if verified is None:
            skipped[reason] = skipped.get(reason, 0) + 1
            continue
        row = make_fields(product, verified, address)
        validate_row(row, product["Группа"])
        rows[product["Группа"]].append(row)

    total = sum(map(len, rows.values()))
    if total == 0:
        raise ValueError("Нет полностью проверенных карточек из свежего прайса. XLSX не формируются.")
    # Проверяем все шаблоны перед записью, чтобы избежать неполного комплекта
    built = {}
    for group_name, data in rows.items():
        if not data:
            continue
        template_rel = text(avito.get(GROUPS[group_name]["template_key"]))
        if not template_rel and not allow_generated_template:
            raise ValueError(
                f"Не указан оригинальный шаблон {group_name}: настройте avito.{GROUPS[group_name]['template_key']} "
                "или используйте --allow-generated-template для тестового файла."
            )
        template = config_path.parent / template_rel if template_rel else None
        built[group_name] = make_book(data, group_name, template)
    output_dir.mkdir(parents=True, exist_ok=True)
    for group_name, book in built.items():
        target = output_dir / GROUPS[group_name]["filename"]
        temporary = target.with_suffix(".tmp.xlsx")
        book.save(temporary)
        temporary.replace(target)
        book.close()
        print(f"{group_name}: {len(rows[group_name])} объявлений → {target}")
    print(f"Итого: {total} проверенных карточек; пропущено: {skipped}")


def main() -> int:
    p = argparse.ArgumentParser(description="РоссКо → Авито: проверенные XLSX Иркутск")
    p.add_argument("--config", default="config.json")
    p.add_argument("--approvals", default="input/avito_approved.csv")
    p.add_argument("--out-dir", default="output/avito_irkutsk")
    p.add_argument("--confirm-irkutsk-source", action="store_true")
    p.add_argument("--allow-generated-template", action="store_true",
                   help="Только для теста; без оригинальных шаблонов нужно проверить XLSX на Авито")
    args = p.parse_args()
    if not args.confirm_irkutsk_source:
        p.error("Подтвердите, что прайс РОССКО относится к Иркутску")
    config_path = Path(args.config).resolve()
    approvals_path = Path(args.approvals)
    if not approvals_path.is_absolute():
        approvals_path = config_path.parent / approvals_path
    target = Path(args.out_dir)
    if not target.is_absolute():
        target = config_path.parent / target
    build(config_path, approvals_path, target, allow_generated_template=args.allow_generated_template)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
