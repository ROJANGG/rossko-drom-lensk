#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Find Rossko public product photos and write input/photos.csv.

The Drom exporter already reads input/photos.csv with columns:
brand;article;photo_url

This script uses the Rossko price list and the "Номенклатура" column to open
public Rossko card pages and extract direct imgs.rossko.ru links.
"""

from __future__ import annotations

import argparse
import csv
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, Optional
from urllib.parse import quote, unquote

import requests
from openpyxl import load_workbook

from update_prices import download_price_if_needed, load_config, normalize_key, text

REQUIRED_HEADERS = ["Номенклатура", "Бренд", "Артикул"]

# Rossko pages may contain both absolute URLs and protocol-relative URLs:
# https://imgs.rossko.ru/15/5A/NSII0012963910/1.jpg
# //imgs.rossko.ru/15/5A/NSII0012963910/1.jpg
DIRECT_IMG_RE = re.compile(
    r"(?:(?:https?:)?//)imgs\.rossko\.ru/[^\"'<>\s)]+?\.(?:jpg|jpeg|png|webp)",
    re.IGNORECASE,
)

DEFAULT_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/123.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
    "Accept-Language": "ru-RU,ru;q=0.9,en;q=0.8",
}

_thread_local = threading.local()


@dataclass(frozen=True)
class PriceRow:
    brand: str
    article: str
    nomenclature: str

    @property
    def key(self) -> str:
        return f"{normalize_key(self.brand)}|{normalize_key(self.article)}"


@dataclass(frozen=True)
class WorkItem:
    index: int
    row: PriceRow


def get_thread_session() -> requests.Session:
    session = getattr(_thread_local, "session", None)
    if session is None:
        session = requests.Session()
        session.headers.update(DEFAULT_HEADERS)
        _thread_local.session = session
    return session


def slugify(value: str) -> str:
    """Build the visible part of Rossko card URL."""
    raw = text(value).casefold()
    raw = re.sub(r"[^0-9a-zа-яё]+", "-", raw, flags=re.IGNORECASE)
    raw = re.sub(r"-+", "-", raw).strip("-")
    return quote(raw)


def card_url(row: PriceRow) -> str:
    slug = "-".join(
        part for part in [slugify(row.brand), slugify(row.article), slugify(row.nomenclature)] if part
    )
    return f"https://rossko.ru/card/{slug}/"


def read_header_map(ws) -> Dict[str, int]:
    result: Dict[str, int] = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        name = text(cell.value)
        if name:
            result[name] = col_idx
    return result


def get_cell(row, header_map: Dict[str, int], header: str):
    idx = header_map.get(header)
    if not idx:
        return None
    return row[idx - 1].value


def iter_price_rows(path: Path) -> Iterable[PriceRow]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    header_map = read_header_map(ws)
    missing = [h for h in REQUIRED_HEADERS if h not in header_map]
    if missing:
        raise ValueError("В прайсе Росско не найдены колонки: " + ", ".join(missing))

    seen: set[str] = set()
    for row in ws.iter_rows(min_row=2):
        item = PriceRow(
            brand=text(get_cell(row, header_map, "Бренд")),
            article=text(get_cell(row, header_map, "Артикул")),
            nomenclature=text(get_cell(row, header_map, "Номенклатура")),
        )
        if not item.brand or not item.article or not item.nomenclature:
            continue
        if item.key in seen:
            continue
        seen.add(item.key)
        yield item


def load_existing_photos(path: Path) -> Dict[str, tuple[str, str, str]]:
    """Return key -> (brand, article, photo_url). Keeps manual rows."""
    result: Dict[str, tuple[str, str, str]] = {}
    if not path.exists():
        return result

    with path.open("r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.DictReader(f, delimiter=delimiter)
        for raw in reader:
            brand = text(raw.get("brand") or raw.get("Бренд"))
            article = text(raw.get("article") or raw.get("Артикул"))
            photo_url = text(raw.get("photo_url") or raw.get("Фотография") or raw.get("url"))
            if brand and article and photo_url:
                result[f"{normalize_key(brand)}|{normalize_key(article)}"] = (brand, article, photo_url)
    return result


def save_photos(path: Path, photos: Dict[str, tuple[str, str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = sorted(photos.values(), key=lambda x: (normalize_key(x[0]), normalize_key(x[1])))
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["brand", "article", "photo_url"])
        writer.writerows(rows)


def decode_html(html: str) -> str:
    # Rossko pages may contain raw URLs, escaped URLs, protocol-relative URLs,
    # or percent-encoded nested image URLs.
    decoded = html.replace("\\/", "/")
    for _ in range(5):
        new_value = unquote(decoded)
        if new_value == decoded:
            break
        decoded = new_value
    return decoded


def clean_image_url(url: str) -> str:
    cleaned = text(url).replace("&amp;", "&").split("?")[0]
    if cleaned.startswith("//"):
        cleaned = "https:" + cleaned
    return cleaned


def image_exists(session: requests.Session, url: str, timeout: int) -> bool:
    try:
        response = session.get(
            url,
            timeout=timeout,
            stream=True,
            allow_redirects=True,
            headers={"Range": "bytes=0-1024", "Referer": "https://rossko.ru/"},
        )
        response.close()
    except requests.RequestException:
        return False
    if response.status_code not in (200, 206):
        return False
    content_type = response.headers.get("content-type", "").lower()
    return (not content_type) or content_type.startswith("image/")


def normalize_to_main_image(session: requests.Session, found_url: str, timeout: int, verify: bool) -> Optional[str]:
    cleaned = clean_image_url(found_url)

    # Fast mode: trust the direct image URL already present on the Rossko card.
    # This avoids 1-9 extra image requests per product.
    if not verify:
        return cleaned

    base = re.sub(r"/\d+\.(?:jpg|jpeg|png|webp)$", "/", cleaned, flags=re.IGNORECASE)

    # Drom should receive the original Rossko image. Prefer 1.jpg, but test nearby
    # numbers because some cards expose another image first.
    candidates = [f"{base}{num}.jpg" for num in (1, 2, 3, 4, 5, 6, 7, 8, 9)]
    if cleaned not in candidates:
        candidates.insert(0, cleaned)

    for candidate in candidates:
        if image_exists(session, candidate, timeout):
            return candidate
    return None


def extract_image_candidates(html: str, row: PriceRow) -> list[str]:
    html = decode_html(html)
    nomenclature_key = f"/{row.nomenclature}/".casefold()
    matches: list[str] = []

    for match in DIRECT_IMG_RE.findall(html):
        clean = clean_image_url(match)
        if nomenclature_key in clean.casefold() and clean not in matches:
            matches.append(clean)

    return matches


def find_photo_url(session: requests.Session, row: PriceRow, timeout: int, verify_image: bool) -> Optional[str]:
    url = card_url(row)
    try:
        response = session.get(url, timeout=timeout, allow_redirects=True)
    except requests.RequestException as exc:
        print(f"HTTP ERROR: {row.brand} {row.article} {url} -> {exc}")
        return None
    if response.status_code >= 400:
        print(f"HTTP {response.status_code}: {row.brand} {row.article} {url}")
        return None

    for match in extract_image_candidates(response.text, row):
        normalized = normalize_to_main_image(session, match, timeout, verify_image)
        if normalized:
            return normalized
    return None


def read_progress(path: Optional[Path]) -> Optional[int]:
    if not path or not path.exists():
        return None
    raw = text(path.read_text(encoding="utf-8", errors="ignore"))
    try:
        value = int(raw)
    except ValueError:
        return None
    return max(0, value)


def save_progress(path: Optional[Path], value: int) -> None:
    if not path:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(str(max(0, value)) + "\n", encoding="utf-8")


def select_work_items(
    price_path: Path,
    photos: Dict[str, tuple[str, str, str]],
    start: int,
    limit: int,
    force: bool,
) -> tuple[list[WorkItem], int, int, Optional[int]]:
    work_items: list[WorkItem] = []
    skipped_existing = 0
    processed_after_start = 0
    last_seen_index: Optional[int] = None

    for idx, row in enumerate(iter_price_rows(price_path)):
        if idx < start:
            continue

        processed_after_start += 1
        last_seen_index = idx

        if not force and row.key in photos:
            skipped_existing += 1
            continue

        if limit and len(work_items) >= limit:
            break

        work_items.append(WorkItem(index=idx, row=row))

    return work_items, skipped_existing, processed_after_start, last_seen_index


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Update input/photos.csv from public Rossko cards")
    parser.add_argument("--config", default="config.json", help="Path to config.json")
    parser.add_argument("--limit", type=int, default=50, help="How many new products to scan. 0 = scan all")
    parser.add_argument("--start", type=int, default=0, help="Skip first N products from price before scanning")
    parser.add_argument("--sleep", type=float, default=0.0, help="Pause between Rossko card request submissions")
    parser.add_argument("--timeout", type=int, default=20, help="HTTP timeout in seconds")
    parser.add_argument("--workers", type=int, default=1, help="Parallel workers. Use 1 for old sequential mode")
    parser.add_argument("--force", action="store_true", help="Re-check products that already have photos")
    parser.add_argument("--no-verify-image", action="store_true", help="Do not make extra request to verify image URL")
    parser.add_argument("--progress-file", default="", help="Write next start position to this file")
    parser.add_argument("--use-progress", action="store_true", help="Read start from --progress-file when it exists")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    config_path = Path(args.config)
    base_dir = config_path.resolve().parent
    config = load_config(config_path)

    progress_file = base_dir / args.progress_file if args.progress_file else None
    progress_start = read_progress(progress_file) if args.use_progress else None
    if progress_start is not None:
        args.start = progress_start

    price_path = download_price_if_needed(config, base_dir)
    photo_map_path = base_dir / text(config.get("drom", {}).get("photo_map_path") or "input/photos.csv")
    photos = load_existing_photos(photo_map_path)

    verify_image = not args.no_verify_image
    workers = max(1, args.workers)

    work_items, skipped_existing, processed_after_start, last_seen_index = select_work_items(
        price_path=price_path,
        photos=photos,
        start=max(0, args.start),
        limit=max(0, args.limit),
        force=args.force,
    )

    found = 0
    scanned = 0

    def process_item(item: WorkItem) -> tuple[WorkItem, Optional[str]]:
        session = get_thread_session()
        return item, find_photo_url(session, item.row, args.timeout, verify_image)

    if workers == 1:
        session = get_thread_session()
        for item in work_items:
            scanned += 1
            photo_url = find_photo_url(session, item.row, args.timeout, verify_image)
            if photo_url:
                photos[item.row.key] = (item.row.brand, item.row.article, photo_url)
                found += 1
                print(f"FOUND {found}/{scanned}: {item.row.brand} {item.row.article} -> {photo_url}")
            else:
                print(f"MISS {scanned}: {item.row.brand} {item.row.article} {item.row.nomenclature}")

            if args.sleep > 0:
                time.sleep(args.sleep)
    else:
        print(f"Parallel scan: workers={workers}, selected={len(work_items)}, verify_image={verify_image}")
        with ThreadPoolExecutor(max_workers=workers) as executor:
            future_map = {}
            for item in work_items:
                future_map[executor.submit(process_item, item)] = item
                if args.sleep > 0:
                    time.sleep(args.sleep)

            for future in as_completed(future_map):
                item = future_map[future]
                scanned += 1
                try:
                    _, photo_url = future.result()
                except Exception as exc:  # noqa: BLE001 - keep workflow resilient
                    photo_url = None
                    print(f"ERROR {scanned}: {item.row.brand} {item.row.article} {item.row.nomenclature} -> {exc}")

                if photo_url:
                    photos[item.row.key] = (item.row.brand, item.row.article, photo_url)
                    found += 1
                    print(f"FOUND {found}/{scanned}: {item.row.brand} {item.row.article} -> {photo_url}")
                else:
                    print(f"MISS {scanned}: {item.row.brand} {item.row.article} {item.row.nomenclature}")

    save_photos(photo_map_path, photos)

    next_start = max(0, args.start)
    if work_items:
        next_start = max(item.index for item in work_items) + 1
    elif processed_after_start > 0 and last_seen_index is not None:
        next_start = last_seen_index + 1
    save_progress(progress_file, next_start)

    print("----")
    print(f"photos.csv: {photo_map_path}")
    print(f"progress file: {progress_file or ''}")
    print(f"start: {args.start}")
    print(f"next start: {next_start}")
    print(f"existing skipped: {skipped_existing}")
    print(f"selected for scan: {len(work_items)}")
    print(f"scanned: {scanned}")
    print(f"found: {found}")
    print(f"total photos in file: {len(photos)}")
    if processed_after_start == 0:
        print("В прайсе не найдено строк для обработки после --start.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
