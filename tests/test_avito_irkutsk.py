import unittest
from avito_irkutsk import category, zero_days, select_rows, load_sku_whitelist
from pathlib import Path
import tempfile
from openpyxl import Workbook


class SelectionTests(unittest.TestCase):
    def test_categories(self):
        self.assertEqual(category("Рейка рулевая"), "Рулевые рейки")
        self.assertEqual(category("Диск тормозной"), "Тормозные диски")
        self.assertIsNone(category("Тормозные колодки"))
        self.assertIsNone(category("Рулевая тяга"))
        self.assertIsNone(category("Сайлентблок рулевой рейки"))
        self.assertIsNone(category("Съемник шарнира рулевой рейки"))
        self.assertIsNone(category("Очиститель тормозных дисков, 650 мл"))
        self.assertIsNone(category("Проставка рулевой рейки"))
        self.assertIsNone(category("Сальник рулевой рейки"))
        self.assertEqual(category("Диск тормозной передний"), "Тормозные диски")
        self.assertEqual(category("Рулевая рейка в сборе"), "Рулевые рейки")

    def test_pickup_address_required(self):
        wb = Workbook()
        wb.active.append(["Бренд", "Артикул", "Описание", "Цена, руб.", "Наличие", "Срок поставки, дн."])
        with self.assertRaisesRegex(ValueError, "адрес выдачи"):
            select_rows(wb.active, {"avito": {}})


    def test_irkutsk_local_stock_selection(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Бренд", "Артикул", "Описание", "Цена, руб.", "Наличие", "Срок поставки, дн.", "Город склада"])
        ws.append(["ATE", "123", "Диск тормозной", 1000, 5, 0, "Иркутск"])
        ws.append(["ATE", "124", "Диск тормозной", 1000, 0, 0, "Иркутск"])
        ws.append(["ATE", "125", "Диск тормозной", 1000, -1, 0, "Иркутск"])
        ws.append(["ATE", "126", "Диск тормозной", 1000, 5, 2, "Иркутск"])
        ws.append(["ATE", "127", "Диск тормозной", 1000, 5, 0, "Новосибирск"])
        ws.append(["ATE", "128", "Тормозные колодки", 1000, 5, 0, "Иркутск"])
        ws.append(["ATE", "129", "Рейка рулевая", 3000, 3, 0, "Иркутск"])
        cfg = {"avito": {"pickup_address": "г. Иркутск, ул. Каховского, д. 11/1"},
               "pricing": {"markup_percent": 37, "round_up_to": 50}}
        rows = select_rows(ws, cfg)
        self.assertEqual({x["Артикул"] for x in rows}, {"123", "129"})
        self.assertEqual(next(x["Цена Авито"] for x in rows if x["Артикул"] == "123"), 1400)
        self.assertTrue(all(x["Город"] == "Иркутск" for x in rows))

    def test_avito_price_tiers_are_applied_to_supplier_stock(self):
        from avito_pricing import progressive_price
        wb = Workbook()
        ws = wb.active
        ws.append(["Бренд", "Артикул", "Описание", "Цена, руб.",
                   "Наличие", "Срок поставки, дн.", "Город склада"])
        ws.append(["CTR", "CHEAP", "Диск тормозной", 3000, 2, 0, "Иркутск"])
        ws.append(["CTR", "EXPENSIVE", "Диск тормозной", 30000, 2, 0, "Иркутск"])
        pricing = {"round_up_to": 50, "tiers": [
            {"up_to": 2000, "rate_percent": 37},
            {"up_to": 5000, "rate_percent": 25},
            {"up_to": 10000, "rate_percent": 18},
            {"up_to": 20000, "rate_percent": 12},
            {"up_to": None, "rate_percent": 9},
        ]}
        config = {"avito": {
            "pickup_address": "г. Иркутск, ул. Каховского, д. 11/1",
            "pricing": pricing,
        }, "pricing": {"markup_percent": 37, "round_up_to": 50}}
        result = {row["Артикул"]: row["Цена Авито"] for row in select_rows(ws, config)}
        self.assertEqual(result, {"CHEAP": 4000, "EXPENSIVE": 34500})
        self.assertEqual(result["EXPENSIVE"], progressive_price(30000, pricing))


    def test_pilot_sku_whitelist(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "pilot.csv"
            path.write_text("brand;article\nKrauf;HNQ2139GQ\nCTR;GD0156\n", encoding="utf-8")
            self.assertEqual(load_sku_whitelist(path), {"krauf|hnq2139gq", "ctr|gd0156"})
            path.write_text("brand;article\nCTR;GD0156\nCTR;GD0156\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "повторяется"):
                load_sku_whitelist(path)

    def test_local_zero(self):
        self.assertTrue(zero_days("0"))
        self.assertTrue(zero_days("0,0"))
        self.assertFalse(zero_days(""))
        self.assertFalse(zero_days("2"))


if __name__ == "__main__":
    unittest.main()
