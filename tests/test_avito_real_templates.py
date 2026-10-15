"""Проверяем генератор на двух НАСТОЯЩИХ XLSX-шаблонах из репозитория."""
import unittest
from pathlib import Path

from openpyxl import load_workbook

from avito_feed import GROUPS, HEADERS, approved_payload, make_book, make_fields

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / "input" / "avito_templates"
FILES = {
    "Рулевые рейки": TEMPLATES / "Для автомобилей - Рулевое управление - Шаблон 13-10-2026.xlsx",
    "Тормозные диски": TEMPLATES / "Для автомобилей - Тормозная система - Шаблон 13-10-2026.xlsx",
}


class RealAvitoTemplatesTest(unittest.TestCase):
    def test_original_templates_present_and_exact(self):
        for group, path in FILES.items():
            with self.subTest(group=group):
                self.assertTrue(path.is_file(), "В GitHub должен быть оригинальный шаблон Авито")
                wb = load_workbook(path, read_only=True, data_only=True)
                try:
                    self.assertEqual(wb.sheetnames, ["Инструкция", "Объявления", "Справочники"])
                    ws = wb["Объявления"]
                    actual = tuple(ws.cell(2, j).value for j in range(1, 72))
                    self.assertEqual(actual, HEADERS)
                    self.assertTrue(ws.cell(1, 1).value.endswith(GROUPS[group]["category"]))
                    # В режиме read_only библиотека может брать неверный max_row из метаданных XLSX.
                    # Наличие пустой первой строки объявлений проверяем напрямую.
                    self.assertIsNone(ws.cell(5, 1).value)
                finally:
                    wb.close()

    def test_one_test_row_fits_both_original_templates(self):
        for group, path in FILES.items():
            with self.subTest(group=group):
                product = {
                    "Группа": group, "Бренд": "TESTBRAND",
                    "Артикул": "XYZ123", "Цена Авито": 1400,
                }
                review = {
                    "approved": "yes", "pickup_confirmed": "yes",
                    "condition_confirmed": "yes",
                    "photo_urls": "https://example.com/example-photo.jpg",
                    "fitment": "Пример: модель и год — проверенная применимость",
                    "oem_number": "",
                }
                payload, reason = approved_payload(product, review)
                self.assertIsNotNone(payload, reason)
                fields = make_fields(product, payload, "г. Иркутск, ул. Каховского, д. 11/1")
                book = make_book([fields], group, path)
                try:
                    self.assertEqual(book.sheetnames, ["Инструкция", "Объявления", "Справочники"])
                    ws = book["Объявления"]
                    self.assertEqual(ws.cell(5, 49).value, 1400)
                    self.assertEqual(ws.cell(5, 55).value, "Новое")
                    self.assertEqual(ws.cell(5, 57).value, "В наличии")
                    self.assertEqual(ws.cell(5, 54).value, GROUPS[group]["category"])
                    self.assertTrue(ws.cell(5, 4).value.startswith("irk-"))
                    self.assertEqual(ws.cell(5, 19).value, "https://example.com/example-photo.jpg")
                    self.assertEqual(ws.cell(5, 1).value, "г. Иркутск, ул. Каховского, д. 11/1")
                finally:
                    book.close()


if __name__ == "__main__":
    unittest.main()
