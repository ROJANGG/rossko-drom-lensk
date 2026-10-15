import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook

from avito_feed import (HEADERS, GROUPS, approved_payload, make_book, make_fields,
                        safe_image_urls, stable_id, validate_row)


class AvitoFeedTests(unittest.TestCase):
    def setUp(self):
        self.product = {"Группа": "Тормозные диски", "Бренд": "TEST",
                        "Артикул": "123ABC", "Цена Авито": 1400}
        self.approval = {
            "approved": "yes", "pickup_confirmed": "yes", "condition_confirmed": "yes",
            "photo_urls": "https://example.com/brake.jpg", "fitment": "Test Car 2010-2012",
            "oem_number": "",
        }

    def test_headers(self):
        self.assertEqual(len(HEADERS), 71)

    def test_approval_required(self):
        self.assertIsNone(approved_payload(self.product, None)[0])
        a = dict(self.approval, fitment="")
        self.assertIsNone(approved_payload(self.product, a)[0])
        a = dict(self.approval, photo_urls="http://example.com/a.jpg")
        self.assertIsNone(approved_payload(self.product, a)[0])
        self.assertIsNotNone(approved_payload(self.product, self.approval)[0])

    def test_stable_ids(self):
        a = stable_id("Тормозные диски", "ATE", "123")
        self.assertEqual(a, stable_id("Тормозные диски", "ate", "123"))
        self.assertNotEqual(a, stable_id("Рулевые рейки", "ATE", "123"))

    def test_xlsx(self):
        verified, _ = approved_payload(self.product, self.approval)
        fields = make_fields(self.product, verified, "г. Иркутск, ул. Каховского, 11/1")
        validate_row(fields, "Тормозные диски")
        book = make_book([fields], "Тормозные диски")
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "test.xlsx"
            book.save(path)
            opened = load_workbook(path, read_only=True)
            ws = opened["Объявления"]
            self.assertEqual(ws.cell(2, 49).value, "Цена")
            self.assertEqual(ws.cell(5, 49).value, 1400)
            self.assertEqual(ws.cell(5, 62).value, "Test Car 2010-2012")
            self.assertTrue(ws.cell(5, 4).value.startswith("irk-"))
            opened.close()
        book.close()


if __name__ == "__main__":
    unittest.main()
